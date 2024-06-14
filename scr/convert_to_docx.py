import pandas as pd
from docx import Document

def excel_to_word(excel_file, word_file):
    # Read the Excel file
    df = pd.read_excel(excel_file)

    # Create a new Word document
    doc = Document()

    # Add a table to the Word document
    table = doc.add_table(rows=1, cols=len(df.columns))

    # Add the column names to the table
    for i, column_name in enumerate(df.columns):
        table.cell(0, i).text = str(column_name)

    # Add the data to the table
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

    # Save the Word document
    doc.save(word_file)

if __name__ == "__main__":
    excel_file = 'data/merged_data_test_task.xlsx' 
    word_file = 'merged_data_test_task.docx' 
    excel_to_word(excel_file, word_file)
